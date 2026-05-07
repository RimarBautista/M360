import os
import string
import time
from datetime import date, datetime

import allure
import cv2
import numpy as np

import pyautogui
from allure_commons.types import AttachmentType
from selenium.webdriver import Keys, ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.support.wait import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import openpyxl
from selenium import webdriver

import m360reports


class SeleniumHelper():
    def __init__(self, driver):
        self.upper_alphabet = None
        self.driver = driver

    def enter_text(self, locator, text):
        WebDriverWait(self.driver, 10).until(EC.visibility_of_element_located(locator)).send_keys(text)

    def capture_screenshot(self,file_name):
        self.driver.get_screenshot_as_file(file_name)

    def click_element(self, locator):
        WebDriverWait(self.driver, 10).until(EC.visibility_of_element_located(locator)).click()

    def get_title_page(self):
        return self.driver.title

    def scroll_into_view(self,locator):
        actions = ActionChains(self.driver)
        actions.scroll_to_element(locator).perform()

    def scroll_down(self,verical_ordinate):
        element = self.driver.find_element(By.XPATH, "//div[@id='content-wrapper-div']")
        for i in range(0, 6):
            self.driver.execute_script("arguments[0].scrollTop = arguments[1]", element, verical_ordinate)
            verical_ordinate += 150
            time.sleep(1)

    def get_val(self, locator):
        return WebDriverWait(self.driver, 10).until(EC.visibility_of_element_located(locator)).text

    def web_element_display(self, locator):
        return WebDriverWait(self.driver, 10).until(EC.visibility_of_element_located(locator))

    def create_header(self):
        self.workbook = openpyxl.load_workbook("ExelFiles/Reports.xlsx")
        self.sheet = self.workbook['Reports']
        for row in range(1, self.sheet.max_row+1):
            header_value = self.sheet.cell(row=row, column=1).value
        if header_value == "Email":
            print("Header existed")
        elif header_value is None:
            header = ["Email",
                      "Total Delivered",
                      "SMS Sent",
                      "SMS Acknowledged",
                      "SMS Failed",
                      "SMS Received",
                      "Date Tested",
                      "Domestic SMS",
                      "Dom SMS Premium Sender ID",
                      "International SMS",
                      "Domestic SMS Sent",
                      "Domestic SMS Acknowledged",
                      "Domestic SMS Failed",
                      "Dom SMS Prem Sender ID Sent",
                      "Dom SMS Prem Sender ID Ack",
                      "Dom SMS Prem Sender ID Failed",
                      "International SMS Sent",
                      "International SMS Ack",
                      "International SMS Failed",
                      "Broadcast Sent",
                      "Broadcast Acknowledged",
                      "Broadcast Failed",
                      "Broadcast Received",
                      "Two Way Message Sent",
                      "Two Way Message Acknowledged",
                      "Two Way Message Failed",
                      "Two Way Message Received",
                      "SMS Poll Sent",
                      "SMS Poll Acknowledged",
                      "SMS Poll Failed",
                      "SMS Poll Received",
                      "API Sent",
                      "API Acknowledged",
                      "API Failed",
                      "API Received"]
            for index, value in enumerate(header, start=1):
                self.sheet.cell(row=1, column=index).value = value
            self.upper_alphabet = list(string.ascii_uppercase)
            # print(self.upper_alphabet)
            for i in range(1,len(self.upper_alphabet)):
                self.workbook.active.column_dimensions[self.upper_alphabet[i]].width = 25
                self.workbook.active.column_dimensions["A"+self.upper_alphabet[i]].width = 25
            self.workbook.save("ExelFiles/Reports.xlsx")

    def insert_value_in_report(self,email,
                               total,
                               SMS_sent,
                               SMS_ack,
                               SMS_failed,
                               SMS_received,
                               Domestic_SMS,
                               Domestic_SMS_Premium_Sender,
                               International_SMS,
                               Domestic_SMS_Sent,
                               Domestic_SMS_Acknowledged,
                               DOMESTIC_SMS_Failed,
                               DOMESTIC_SMS_Premium_ID_Sender,
                               domestic_SMS_premium_sender_ID_ack,
                               domestic_SMS_premium_sender_ID_failed,
                               international_SMS_sent,
                               international_SMS_ack,
                               international_SMS_failed,
                               broadcast_sent,
                               broadcast_ack,
                               broadcast_failed,
                               broadcast_received,
                               two_way_message_sent,
                               two_way_message_ack,
                               two_way_message_failed,
                               two_way_message_received,
                               SMS_poll_sent,
                               SMS_poll_ack,
                               SMS_poll_failed,
                               SMS_poll_received,
                               api_sent,
                               api_ack,
                               api_failed,
                               api_received
                               ):
        self.workbook = openpyxl.load_workbook("ExelFiles/Reports.xlsx")
        self.sheet = self.workbook['Reports']
        count = 2
        for col in range(1, self.sheet.max_row + 1):
            col_value = self.sheet.cell(row=count, column=col).value
            print(col_value)
            if col_value is not None:
                count += 1
                print("Entered in not none")
                print(count)
            else:
                print(col_value)
                data = [email,total,
                        SMS_sent,
                        SMS_ack,
                        SMS_failed,
                        SMS_received,
                        datetime.now(),
                        Domestic_SMS,
                        Domestic_SMS_Premium_Sender,
                        International_SMS,
                        Domestic_SMS_Sent,
                        Domestic_SMS_Acknowledged,
                        DOMESTIC_SMS_Failed,
                        DOMESTIC_SMS_Premium_ID_Sender,
                        domestic_SMS_premium_sender_ID_ack,
                        domestic_SMS_premium_sender_ID_failed,
                        international_SMS_sent,
                        international_SMS_ack,
                        international_SMS_failed,
                        broadcast_sent,
                        broadcast_ack,
                        broadcast_failed,
                        broadcast_received,
                        two_way_message_sent,
                        two_way_message_ack,
                        two_way_message_failed,
                        two_way_message_received,
                        SMS_poll_sent,
                        SMS_poll_ack,
                        SMS_poll_failed,
                        SMS_poll_received,
                        api_sent,
                        api_ack,
                        api_failed,
                        api_received
                        ]
                for index, value in enumerate(data, start=1):
                    self.sheet.cell(row=count, column=index).value = value
                self.workbook.save("ExelFiles/Reports.xlsx")
                break

    @staticmethod
    def start_video_recording():
        path = "m360reports/"
        os.makedirs(path, exist_ok=True)
        screen_width, screen_height = pyautogui.size()
        print("Screen resolution size: {}x{}".format(screen_width, screen_height))
        print("Video Recording Started")
        global video_recording
        video_recording = True
        resolution = (1920, 1080)
        codec = cv2.VideoWriter_fourcc(*"mp4v")
        fps = 60.0
        timestamp = datetime.now().strftime("%d%m%Y_%H%M%S")
        filename_video = os.path.join(f"video_{timestamp}.MP4")
        out = cv2.VideoWriter(filename_video, codec, fps, resolution)
        while video_recording:
            img = pyautogui.screenshot()
            frame = np.array(img)
            frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            out.write(frame)
            if cv2.waitKey(1) == ord('q'):
                break
        out.release()
        if filename_video is not None:
            allure.attach.file(filename_video, name="Test_Video", attachment_type=AttachmentType.MP4)
            print(f"Video attached: {filename_video}")

    @staticmethod
    def stop_video_recording():
        global video_recording
        video_recording = False
        cv2.destroyAllWindows()
        print("Video Recording Stopped.")
